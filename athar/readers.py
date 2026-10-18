"""Readers for the two Excel inputs.

1) Odoo POS pivot export ("Point of Sale Analysis", measure "Subtotal w/o discount"),
   hierarchy by leading spaces in column A: 5 = company, 10 = branch, 15 = POS user.
   Two layouts are supported:
     (a) one value column (month-to-date total)
     (b) one column per day ("01 Oct 2026", "02 Oct 2026", ...) + a final "Total" column  <- preferred
   Values are in each country's own currency. Odoo's grand "Total" row mixes currencies: ignored.

2) Monthly target file: columns  #, Country, Region, City, Store, Target  (target in local currency).
"""
import re
import datetime
import openpyxl

from .config import EXCLUDE

_DAY = re.compile(r"^(\d{2}) (\w{3}) (\d{4})$")


def read_pos(path):
    """Return (stores, daily).

    stores = {CODE: {"sales": month_to_date, "users": [[user, sales], ...]}}
    daily  = {"YYYY-MM-DD": {CODE: sales_that_day}}   (empty for layout (a))
    """
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    dates = []
    for r in rows[:4]:
        found = [(i, c) for i, c in enumerate(r) if isinstance(c, str) and _DAY.match(c)]
        if len(found) >= 2:
            dates = [(i, datetime.datetime.strptime(c, "%d %b %Y").date().isoformat()) for i, c in found]
            break
    vcol = (dates[-1][0] + 1) if dates else 1          # Total column
    stores, daily, store = {}, {d: {} for _, d in dates}, None
    for r in rows[2:]:
        name = r[0]
        if not isinstance(name, str):
            continue
        indent = len(name) - len(name.lstrip(" "))
        n = name.strip()
        if indent == 10:
            store = n.upper()
            stores[store] = {"sales": round(float(r[vcol] or 0), 2), "users": []}
            for i, d in dates:
                daily[d][store] = round(float(r[i] or 0), 2)
        elif indent == 15 and store:
            stores[store]["users"].append([n, round(float(r[vcol] or 0), 2)])
    if not stores:
        raise ValueError("No branches found — is this the Odoo pivot export (company > branch > user)?")
    for code in EXCLUDE:
        stores.pop(code, None)
        for d in daily:
            daily[d].pop(code, None)
    return stores, daily


def read_targets(path):
    """Return {CODE: {"country", "region", "city", "target"}} (target in local currency)."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[4] is None:
            continue
        out[str(r[4]).strip().upper()] = {"country": r[1], "region": r[2], "city": r[3],
                                         "target": round(float(r[5] or 0), 2)}
    if not out:
        raise ValueError("Target file is empty or has different columns (#, Country, Region, City, Store, Target)")
    return out
